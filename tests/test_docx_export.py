import tempfile
import unittest
from pathlib import Path

from docx import Document

from backend.api.routes import _demote_body_markdown_headings, _numbered_export_sections, _strip_duplicate_section_heading
from backend.export.md_to_word import convert_md_to_word


class DocxExportRegressionTest(unittest.TestCase):
    def test_body_markdown_headings_are_not_exported_as_word_outline_headings(self):
        content = "\n".join(
            [
                "## 噪声与振动控制措施",
                "",
                "### 主要控制措施",
                "",
                "正文内容。",
                "",
                "监测与记录",
                "---",
                "",
                "| 项目 | 内容 |",
                "| --- | --- |",
                "| 监测 | 按要求执行 |",
            ]
        )

        normalized = _demote_body_markdown_headings(content)

        self.assertNotIn("## 噪声与振动控制措施", normalized)
        self.assertNotIn("### 主要控制措施", normalized)
        self.assertIn("【噪声与振动控制措施】", normalized)
        self.assertIn("【主要控制措施】", normalized)
        self.assertIn("【监测与记录】", normalized)
        self.assertIn("| 项目 | 内容 |", normalized)

    def test_docx_navigation_headings_remain_official_section_headings_only(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "outline-lock.md"
            body = _demote_body_markdown_headings(
                "\n".join(
                    [
                        "## 噪声与振动控制措施",
                        "",
                        "### 主要控制措施",
                        "",
                        "正文内容。",
                    ]
                )
            )
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 测试投标文件",
                        "",
                        "## 1. 企业营业执照",
                        "",
                        body,
                        "",
                        "## 2. 安全生产许可证",
                        "",
                        "正文内容。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            headings = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style and paragraph.style.name.startswith("Heading")
            ]

            self.assertIn("1. 企业营业执照", headings)
            self.assertIn("2. 安全生产许可证", headings)
            self.assertNotIn("噪声与振动控制措施", headings)
            self.assertNotIn("主要控制措施", headings)

    def test_duplicate_numbered_body_heading_is_removed_before_docx_export(self):
        content = "## 1. 企业营业执照\n\n正文内容。"
        cleaned = _strip_duplicate_section_heading(content, {"title": "企业营业执照", "level": 1})

        self.assertEqual(cleaned, "正文内容。")

    def test_export_sections_are_numbered_for_word_outline(self):
        sections = [
            {"id": "a", "level": 1, "title": "企业营业执照"},
            {"id": "b", "level": 1, "title": "2. 安全生产许可证"},
            {"id": "c", "level": 1, "title": "项目经理资格"},
            {"id": "d", "level": 2, "title": "项目经理简历表"},
        ]

        numbered = _numbered_export_sections(sections)

        self.assertEqual(numbered[0]["_export_title"], "1. 企业营业执照")
        self.assertEqual(numbered[1]["_export_title"], "2. 安全生产许可证")
        self.assertEqual(numbered[2]["_export_title"], "3. 项目经理资格")
        self.assertEqual(numbered[3]["_export_title"], "3.1 项目经理简历表")

    def test_docx_first_page_is_clickable_toc_and_title_is_not_outline_heading(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "toc.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 招标文件",
                        "",
                        "# 1. 企业营业执照",
                        "",
                        "正文内容。",
                        "",
                        "# 2. 安全生产许可证",
                        "",
                        "正文内容。",
                        "",
                        "## 2.1 安全生产许可范围",
                        "",
                        "正文内容。",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            non_empty_paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            headings = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.style and paragraph.style.name.startswith("Heading")
            ]
            hyperlinks = document._element.xpath(".//w:hyperlink")
            bookmarks = document._element.xpath(".//w:bookmarkStart")

            self.assertEqual(non_empty_paragraphs[0], "招标文件")
            self.assertEqual(non_empty_paragraphs[1], "目录")
            self.assertIn("1. 企业营业执照", non_empty_paragraphs[:8])
            self.assertIn("2. 安全生产许可证", non_empty_paragraphs[:8])
            self.assertNotIn("招标文件", headings)
            self.assertIn("1. 企业营业执照", headings)
            self.assertIn("2. 安全生产许可证", headings)
            self.assertGreaterEqual(len(hyperlinks), 3)
            self.assertGreaterEqual(len(bookmarks), 3)

    def test_formal_docx_cleans_generation_notes_emoji_and_preserves_table(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            markdown_path = Path(tmpdir) / "formal.md"
            markdown_path.write_text(
                "\n".join(
                    [
                        "# 某水库除险加固工程投标文件",
                        "",
                        "## 商务响应文件",
                        "",
                        "（本章节正文共计约2980字，符合目标字数要求，结构完整，可用于直接插入标书资格文件分册）",
                        "",
                        "关键提醒：📌 本投标单位承诺严格响应招标文件要求。",
                        "",
                        "| 项目 | 响应 |",
                        "| --- | --- |",
                        "| 工期 | 满足招标文件要求 |",
                    ]
                ),
                encoding="utf-8",
            )

            output_path = convert_md_to_word(markdown_path)
            document = Document(str(output_path))
            full_text = "\n".join(paragraph.text for paragraph in document.paragraphs)

            self.assertIn("商务响应文件", full_text)
            self.assertIn("本投标单位承诺严格响应招标文件要求", full_text)
            self.assertNotIn("本章节正文共计", full_text)
            self.assertNotIn("📌", full_text)
            self.assertEqual(len(document.tables), 1)
            self.assertEqual(document.tables[0].cell(1, 0).text, "工期")

    def test_markdown_image_limit_records_skipped_images(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp = Path(tmpdir)
            image_path = tmp / "tiny.png"
            image_path.write_bytes(
                bytes.fromhex(
                    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
                    "0000000c49444154789c63606060000000040001f61738550000000049454e44ae426082"
                )
            )
            markdown_path = tmp / "images.md"
            markdown_path.write_text(
                f"# 图文测试\n\n![图1]({image_path})\n\n![图2]({image_path})\n",
                encoding="utf-8",
            )

            import backend.export.md_to_word as md_to_word

            original_limit = md_to_word.MARKDOWN_IMAGE_MAX_COUNT
            md_to_word.MARKDOWN_IMAGE_MAX_COUNT = 1
            try:
                _, report = convert_md_to_word(markdown_path, return_report=True)
            finally:
                md_to_word.MARKDOWN_IMAGE_MAX_COUNT = original_limit

            self.assertEqual(report["found"], 2)
            self.assertEqual(report["inserted"], 1)
            self.assertEqual(report["skipped"], 1)


if __name__ == "__main__":
    unittest.main()
